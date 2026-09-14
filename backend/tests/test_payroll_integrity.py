"""Local in-memory acceptance scenarios. Never connects to the configured database."""
import os
os.environ['DATABASE_URL'] = 'sqlite://'
import unittest
from unittest.mock import patch
from datetime import date, timedelta
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from app.models.models import (Base, Driver, DriverProfile, Load, LoadStatus, BillingStatus,
    Settlement, SettlementStatus, SettlementPayment, SettlementAdjustment, AdvancedPayment,
    LoadService, ServiceType, Expense, PayrollCarryover, LoadStop, StopType, Payment, DriverScheduledTransaction)
from app.schemas.schemas import LoadCreate, LoadUpdate, LoadServiceCreate
from app.schemas.payroll_schemas import SettlementCreate, SettlementPaymentCreate, SettlementAdjustmentCreate
from app.crud import payroll, loads, reports
from app.services.driver_pay_service import take_snapshot, compute_driver_pay, is_locked
from app.api.v1.endpoints import payroll as payroll_api
from app.api.v1.endpoints.drivers_extended import DriverProfileIn, update_driver_extended

TODAY = date(2026, 9, 14)

class PayrollScenarios(unittest.TestCase):
    def setUp(self):
        from sqlalchemy.pool import StaticPool
        self.engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
        Base.metadata.create_all(self.engine)
        self.db = Session(self.engine, autoflush=False)
        self.driver = Driver(name='Driver A', pay_rate_loaded=.60, pay_rate_empty=0, is_active=True)
        self.db.add(self.driver); self.db.flush()
        self.profile = DriverProfile(driver_id=self.driver.id, pay_type='per_mile', payable_to='Driver A', freight_percentage=30)
        self.db.add(self.profile); self.db.commit()

    def tearDown(self):
        self.db.close(); self.engine.dispose()

    def load(self, **kw):
        n = self.db.query(Load).count() + 1001
        data = dict(load_number=n, driver_id=self.driver.id, load_date=TODAY, rate=6000,
                    loaded_miles=2000, empty_miles=100, total_miles=2100, is_active=True,
                    status=LoadStatus.NEW, billing_status=BillingStatus.PENDING)
        data.update(kw)
        l = Load(**data); self.db.add(l); self.db.flush(); take_snapshot(self.db, l); self.db.commit()
        return l

    def settlement(self):
        return payroll.create_settlement(self.db, SettlementCreate(driver_id=self.driver.id, date=TODAY))

    def ready(self, s):
        s.status = SettlementStatus.READY; self.db.commit()

    def schedule(self, **kw):
        data = dict(driver_id=self.driver.id, trans_type='deduction', amount=100,
                    schedule='weekly', start_date=date(2026, 9, 1), repeat_type='times', repeat_times=3)
        data.update(kw)
        tx = DriverScheduledTransaction(**data)
        self.db.add(tx); self.db.commit()
        return tx

    def apply_schedule(self, s, tx, due):
        return payroll_api.apply_scheduled(s.id, tx.id, payroll_api.ScheduledApplicationIn(due_date=due), self.db)

    def test_saved_ezloads_preview_fixtures(self):
        import json
        from pathlib import Path
        from app.services.scheduled_payroll import schedule_dates
        fixtures = json.loads((Path(__file__).parent / 'fixtures' / 'ezloads_schedule_previews.json').read_text())
        for fixture in fixtures:
            with self.subTest(case=fixture['case']):
                req = fixture['request']
                tx = self.schedule(schedule=req['schedule']['type'], start_date=date.fromisoformat(req['start_at']),
                    repeat_type=req['repeat']['type'], repeat_times=req['repeat'].get('times'),
                    end_date=date.fromisoformat(req['repeat']['until']) if 'until' in req['repeat'] else None)
                dates = schedule_dates(tx, date(2027, 1, 1))
                self.assertEqual([d.isoformat() for d in dates], fixture['expected']['dates'])
                self.assertEqual([-tx.amount for _ in dates], fixture['expected']['amounts'])

    def test_reference_calendar_previews(self):
        from app.services.scheduled_payroll import schedule_dates
        tx = self.schedule()
        self.assertEqual(schedule_dates(tx, date(2026, 10, 1)), [date(2026, 9, d) for d in (1, 8, 15)])
        tx.schedule = 'biweekly'
        self.assertEqual(schedule_dates(tx, date(2026, 10, 1)), [date(2026, 9, d) for d in (1, 15, 29)])
        tx.schedule = 'monthly'; tx.start_date = date(2026, 1, 31)
        self.assertEqual(schedule_dates(tx, date(2026, 4, 1)), [date(2026, 1, 31), date(2026, 2, 28), date(2026, 3, 31)])
        tx.schedule = 'weekly'; tx.start_date = date(2026, 9, 1); tx.repeat_type = 'until'; tx.end_date = date(2026, 9, 15)
        self.assertEqual(schedule_dates(tx, date(2026, 10, 1)), [date(2026, 9, d) for d in (1, 8, 15)])

    def test_scheduled_occurrence_once_and_reversible(self):
        from app.services.scheduled_payroll import available_dates
        from fastapi import HTTPException
        tx = self.schedule(); s = self.settlement(); due = date(2026, 9, 1)
        first = self.apply_schedule(s, tx, due)
        again = self.apply_schedule(s, tx, due)
        self.assertEqual(first['id'], again['id'])
        self.assertEqual(s.settlement_total, -100)
        self.assertEqual(tx.times_applied, 1)
        other = self.settlement()
        with self.assertRaises(HTTPException): self.apply_schedule(other, tx, due)
        self.assertNotIn(due, available_dates(self.db, tx, TODAY))
        payroll.delete_adjustment(self.db, s.id, first['id'])
        self.assertEqual(tx.times_applied, 0)
        self.assertIn(due, available_dates(self.db, tx, TODAY))
        self.apply_schedule(other, tx, due)
        payroll.delete_settlement(self.db, other.id)
        self.assertIn(due, available_dates(self.db, tx, TODAY))

    def test_scheduled_rejects_future_paused_wrong_payee_and_legacy(self):
        from fastapi import HTTPException
        tx = self.schedule(); s = self.settlement()
        with self.assertRaises(HTTPException): self.apply_schedule(s, tx, date(2026, 9, 15))
        tx.is_active = False; self.db.commit()
        with self.assertRaises(HTTPException): self.apply_schedule(s, tx, date(2026, 9, 1))
        tx.is_active = True; tx.payable_to = 'Other'; self.db.commit()
        with self.assertRaises(HTTPException): self.apply_schedule(s, tx, date(2026, 9, 1))
        tx.payable_to = None; tx.times_applied = 1; self.db.commit()
        with self.assertRaises(HTTPException): self.apply_schedule(s, tx, date(2026, 9, 1))

    def test_flatpay_persists_separate_schedule_and_preserves_zero_mileage(self):
        from app.api.v1.endpoints.drivers_extended import create_driver_extended
        result = create_driver_extended(DriverProfileIn(name='Flat Driver', pay_type='flatpay', flatpay=500,
            flatpay_period='weekly', flatpay_start_date=date(2026, 9, 1), hire_date=date(2020, 1, 1),
            per_extra_stop=25, pay_rate_loaded=0, pay_rate_empty=0), self.db)
        tx = self.db.query(DriverScheduledTransaction).filter_by(source_key=f"flatpay:{result['id']}").one()
        self.assertEqual(tx.amount, 500)
        self.assertEqual(tx.start_date, date(2026, 9, 1))
        self.assertEqual(result['profile']['hire_date'], '2020-01-01')
        self.assertEqual(result['profile']['per_extra_stop'], 25)
        self.assertEqual(result['pay_rate_loaded'], 0)
        update_driver_extended(result['id'], DriverProfileIn(phone='123'), self.db)
        self.assertEqual(self.db.query(DriverScheduledTransaction).filter_by(driver_id=result['id']).count(), 1)

    def test_hourly_work_freezes_rate_and_applies_only_once(self):
        from fastapi import HTTPException
        self.profile.pay_type = 'hourly'; self.profile.hourly_rate = 25; self.db.commit()
        s = self.settlement()
        data = payroll_api.TimeReportIn(date=TODAY, hours=7.5, description='Warehouse', request_key='hours-fixture-00001')
        report = payroll_api.create_time_report(s.id, data, self.db)
        self.assertEqual(report['amount'], 187.5)
        self.assertEqual(payroll_api.create_time_report(s.id, data, self.db)['id'], report['id'])
        self.profile.hourly_rate = 50; self.db.commit()
        adj = payroll_api.apply_time_report(s.id, report['id'], self.db)
        self.assertEqual(s.settlement_total, 187.5)
        self.assertEqual(payroll_api.apply_time_report(s.id, report['id'], self.db)['id'], adj['id'])
        other = self.settlement()
        with self.assertRaises(HTTPException): payroll_api.apply_time_report(other.id, report['id'], self.db)
        with self.assertRaises(HTTPException): payroll_api.delete_time_report(s.id, report['id'], self.db)
        payroll.delete_adjustment(self.db, s.id, adj['id'])
        self.assertEqual(len(payroll_api.available_time_reports(s.id, self.db)['rows']), 1)
        payroll_api.apply_time_report(other.id, report['id'], self.db)
        payroll.delete_settlement(self.db, other.id)
        self.assertEqual(len(payroll_api.available_time_reports(s.id, self.db)['rows']), 1)

    def test_reference_loan_and_escrow_installments(self):
        import json
        from pathlib import Path
        from app.services.scheduled_payroll import schedule_dates, amount_for_date
        fixtures = json.loads((Path(__file__).parent / 'fixtures' / 'ezloads_loan_previews.json').read_text())
        for fixture in fixtures:
            tx = self.schedule(trans_type='loan' if fixture['type'] == 3 else 'escrow', amount=1000, deduct_by=300, repeat_type='always')
            dates = schedule_dates(tx, date(2026, 12, 1))
            self.assertEqual([d.isoformat() for d in dates], fixture['dates'])
            self.assertEqual([-amount_for_date(tx, d) for d in dates], fixture['amounts'])
            s = self.settlement()
            row = self.apply_schedule(s, tx, dates[0])
            self.assertEqual(row['amount'], 300)
            self.assertEqual(s.settlement_total, -300)
            payroll.delete_adjustment(self.db, s.id, row['id'])
            self.assertEqual(tx.times_applied, 0)
            self.assertEqual(amount_for_date(tx, dates[-1]), 100)

    @patch("app.services.report_export_service._get_company_info", return_value={"name": "KARVAN Test", "email": "", "phone": "", "address": ""})
    def test_report_groups_use_ids_weighted_rates_and_export_totals(self, company):
        from app.services.report_export_service import generate_total_revenue_xlsx, generate_rate_per_mile_pdf
        from openpyxl import load_workbook
        from pypdf import PdfReader
        from io import BytesIO
        self.load(rate=1000, loaded_miles=100, total_miles=100)
        self.load(rate=1000, loaded_miles=900, total_miles=900)
        data = reports.get_total_revenue_report(self.db, date_from=TODAY, date_to=TODAY, group_by='driver')
        self.assertEqual(len(data['groups']), 1)
        self.assertEqual(data['groups'][0]['summary']['rate_per_mile'], 2)
        wb = load_workbook(BytesIO(generate_total_revenue_xlsx(data, {}, [])))
        self.assertEqual(wb['Group summary']['C2'].value, 2000)
        self.assertEqual(wb['Group summary']['E2'].value, 2)
        sheet = wb.worksheets[0]
        self.assertEqual(sheet.cell(sheet.max_row, 7).value, 2000)
        text = ''.join(page.extract_text() for page in PdfReader(BytesIO(generate_rate_per_mile_pdf(data, {}))).pages)
        self.assertIn('Driver A', text)
        self.assertIn('2,000', text)
        another = Driver(name='Driver A', is_active=True); self.db.add(another); self.db.flush()
        self.load(driver_id=another.id)
        data = reports.get_total_revenue_report(self.db, date_from=TODAY, date_to=TODAY, group_by='driver')
        self.assertEqual(len(data['groups']), 2)

    def test_load_override_modes_and_restore_use_frozen_rates(self):
        from app.api.v1.endpoints.loads import set_driver_pay_override, DriverPayOverrideIn
        from fastapi import HTTPException
        load = self.load()
        set_driver_pay_override(load.id, DriverPayOverrideIn(type='fixed', amount=0), self.db)
        self.assertEqual(compute_driver_pay(load), 0)
        set_driver_pay_override(load.id, DriverPayOverrideIn(type='percentage', base=5000, percentage=20), self.db)
        self.assertEqual(compute_driver_pay(load), 1000)
        set_driver_pay_override(load.id, DriverPayOverrideIn(type='per_mile', loaded_rate=.5, empty_rate=0), self.db)
        self.assertEqual(compute_driver_pay(load), 1000)
        self.driver.pay_rate_loaded = 3; self.db.commit()
        set_driver_pay_override(load.id, DriverPayOverrideIn(), self.db)
        self.assertEqual(compute_driver_pay(load), 1200)
        s = self.settlement(); payroll.add_load_item(self.db, s.id, load.id)
        with self.assertRaises(HTTPException): set_driver_pay_override(load.id, DriverPayOverrideIn(type='fixed', amount=1), self.db)

    def test_generated_schedules_survive_pause_and_do_not_duplicate(self):
        from app.services.scheduled_payroll import run_schedules, unassigned_occurrences
        from app.models.models import ScheduledPayrollOccurrence
        tx = self.schedule()
        self.assertEqual(run_schedules(self.db, TODAY), 2)
        self.db.commit()
        self.assertEqual(run_schedules(self.db, TODAY), 0)
        tx.is_active = False; self.db.commit()
        s = self.settlement()
        self.assertEqual(len(unassigned_occurrences(self.db, s)), 2)
        row = self.apply_schedule(s, tx, date(2026, 9, 1))
        self.assertEqual(row['amount'], 100)
        self.assertEqual(self.db.query(ScheduledPayrollOccurrence).count(), 2)
        payroll.delete_adjustment(self.db, s.id, row['id'])
        self.assertEqual(len(unassigned_occurrences(self.db, s)), 2)

    def test_open_balance_includes_generated_pay_and_hourly_work(self):
        self.schedule(trans_type='addition', amount=500)
        self.profile.pay_type = 'hourly'; self.profile.hourly_rate = 20; self.db.commit()
        s = self.settlement()
        payroll_api.create_time_report(s.id, payroll_api.TimeReportIn(date=TODAY, hours=2, request_key='balance-hours-fixture'), self.db)
        balances = payroll.get_open_balances(self.db, date_from=date(2026, 9, 1), date_to=TODAY)
        self.assertEqual(balances[0]['balance'], 1040)
        self.assertEqual(balances[0]['load_ids'], [])

    def test_flatpay_revision_keeps_generated_old_amount(self):
        from app.services.flatpay import sync_flatpay
        from app.services.scheduled_payroll import run_schedules
        from app.models.models import ScheduledPayrollOccurrence
        self.profile.pay_type = 'flatpay'; self.profile.flatpay = 500
        self.profile.flatpay_period = 'weekly'; self.profile.flatpay_start_date = date(2026, 9, 1)
        sync_flatpay(self.db, self.driver, self.profile); self.db.commit()
        run_schedules(self.db, TODAY); self.db.commit()
        self.profile.flatpay = 600; self.profile.flatpay_start_date = date(2026, 9, 15)
        sync_flatpay(self.db, self.driver, self.profile); self.db.commit()
        rows = self.db.query(ScheduledPayrollOccurrence).all()
        self.assertEqual([r.amount for r in rows], [500, 500])
        self.assertEqual(self.db.query(DriverScheduledTransaction).filter_by(source_key=f'flatpay:{self.driver.id}').one().amount, 600)

    def test_zero_rate_is_preserved(self):
        l = self.load()
        self.assertEqual(l.pay_rate_empty_snapshot, 0)
        self.assertEqual(compute_driver_pay(l), 1200)

    def test_profile_edit_does_not_change_existing_pay(self):
        l = self.load()
        self.driver.pay_rate_loaded = 2; self.db.commit()
        self.assertEqual(reports._driver_pay(l), 1200)
        self.assertEqual(compute_driver_pay(l), 1200)

    def test_missing_snapshot_never_uses_live_profile(self):
        l = self.load(); l.drivers_payable_snapshot = None; l.pay_type_snapshot = None
        self.assertEqual(reports._driver_pay(l), 0)

    def test_profile_partial_update_preserves_pay_rules(self):
        self.profile.pay_type = 'percentage'; self.profile.freight_percentage = 35; self.db.commit()
        update_driver_extended(self.driver.id, DriverProfileIn(phone='123'), self.db)
        self.assertEqual(self.profile.pay_type, 'percentage')
        self.assertEqual(self.profile.freight_percentage, 35)

    def test_mileage_change_updates_all_financial_views(self):
        l = self.load()
        loads.update_load(self.db, l.id, LoadUpdate(loaded_miles=3000))
        self.assertEqual(reports._driver_pay(l), 1800)
        self.assertEqual(payroll.get_open_balances(self.db)[0]['balance'], 1800)
        s = self.settlement(); item = payroll.add_load_item(self.db, s.id, l.id)
        self.assertEqual(item.amount, 1800)

    def test_services_apply_to_mileage_pay_and_removal(self):
        l = self.load()
        svc = loads.add_service(self.db, l.id, LoadServiceCreate(service_type=ServiceType.LUMPER,
                                add_deduct='Add', invoice_amount=150, drivers_payable=100))
        self.assertEqual(l.drivers_payable_snapshot, 1300)
        loads.delete_service(self.db, svc.id, l.id)
        self.assertEqual(l.drivers_payable_snapshot, 1200)

    def test_period_pay_is_not_multiplied_by_load_count(self):
        for kind in ['flatpay', 'hourly']:
            self.profile.pay_type=kind; self.profile.flatpay=1500; self.profile.hourly_rate=25; self.db.commit()
            self.assertEqual(self.load().drivers_payable_snapshot, 0)

    def test_enum_lock_is_enforced_before_mutation(self):
        l = self.load(status=LoadStatus.DELIVERED)
        self.assertTrue(is_locked(l))
        with self.assertRaises(ValueError):
            loads.update_load(self.db, l.id, LoadUpdate(status=LoadStatus.NEW, loaded_miles=1))
        self.assertEqual(l.status, LoadStatus.DELIVERED)
        self.assertEqual(l.loaded_miles, 2000)

    def test_settled_load_cannot_change_pay_or_services(self):
        l = self.load(); s = self.settlement(); payroll.add_load_item(self.db,s.id,l.id)
        with self.assertRaises(ValueError):loads.update_load(self.db,l.id,LoadUpdate(rate=7000))
        with self.assertRaises(ValueError):loads.add_service(self.db,l.id,LoadServiceCreate(service_type=ServiceType.LUMPER,drivers_payable=100))

    def test_duplicate_load_in_another_settlement_is_rejected(self):
        l=self.load(); a=self.settlement(); b=self.settlement(); payroll.add_load_item(self.db,a.id,l.id)
        with self.assertRaises(payroll.PayrollError):payroll.add_load_item(self.db,b.id,l.id)

    def test_wrong_driver_load_is_rejected(self):
        other=Driver(name='Other'); self.db.add(other); self.db.commit()
        l=self.load(driver_id=other.id); s=self.settlement()
        with self.assertRaises(payroll.PayrollError):payroll.add_load_item(self.db,s.id,l.id)

    def test_deleted_settlement_releases_load(self):
        l=self.load(); s=self.settlement(); payroll.add_load_item(self.db,s.id,l.id)
        payroll.delete_settlement(self.db,s.id)
        self.assertEqual(payroll.get_open_balances(self.db)[0]['balance'],1200)
        new=self.settlement(); payroll.add_load_item(self.db,new.id,l.id)

    def test_advance_reduces_due_not_total_and_can_be_restored(self):
        l=self.load(); s=self.settlement(); payroll.add_load_item(self.db,s.id,l.id)
        ap=AdvancedPayment(payment_number=3001,driver_id=self.driver.id,payment_date=TODAY,amount=300,applied_amount=0)
        self.db.add(ap); self.db.commit()
        result=payroll_api.apply_advanced_payment(s.id,ap.id,None,self.db)
        self.assertEqual((s.settlement_total,s.balance_due),(1200,900))
        self.assertEqual(ap.applied_amount,300)
        self.assertEqual(self.db.get(SettlementAdjustment,result['id']).advanced_payment_id,ap.id)
        payroll.delete_adjustment(self.db,s.id,result['id'])
        self.assertEqual((s.settlement_total,s.balance_due),(1200,1200))
        self.assertEqual(ap.applied_amount,0)

    def test_delete_settlement_restores_advance(self):
        s=self.settlement();ap=AdvancedPayment(payment_number=3001,driver_id=self.driver.id,payment_date=TODAY,amount=300,applied_amount=0)
        self.db.add(ap);self.db.commit();payroll_api.apply_advanced_payment(s.id,ap.id,None,self.db)
        payroll.delete_settlement(self.db,s.id)
        self.assertEqual(ap.applied_amount,0)

    def test_payment_closes_only_remaining_due(self):
        l=self.load();s=self.settlement();payroll.add_load_item(self.db,s.id,l.id);self.ready(s)
        with self.assertRaises(payroll.PayrollError):payroll.add_payment(self.db,s.id,SettlementPaymentCreate(amount=1201))
        payroll.add_payment(self.db,s.id,SettlementPaymentCreate(amount=1200))
        self.assertEqual((s.settlement_total,s.balance_due,s.status),(1200,0,SettlementStatus.PAID))

    def test_carryover_is_balanced_debt_transfer_not_cash(self):
        s=self.settlement();payroll.add_adjustment(self.db,s.id,SettlementAdjustmentCreate(adj_type='deduction',amount=200))
        self.ready(s);c=payroll.create_carryover(self.db,s.id,TODAY)
        self.assertEqual(s.balance_due,0)
        self.assertEqual(self.db.query(SettlementPayment).count(),0)
        next_s=self.settlement();self.assertEqual(next_s.balance_due,-200)
        self.assertEqual(payroll.create_carryover(self.db,s.id,TODAY).id,c.id)
        self.assertEqual(self.db.query(PayrollCarryover).count(),1)
        third=self.settlement();self.assertEqual(third.balance_due,0)

    def test_positive_balance_cannot_carry_over(self):
        l=self.load();s=self.settlement();payroll.add_load_item(self.db,s.id,l.id);self.ready(s)
        with self.assertRaises(payroll.PayrollError):payroll.create_carryover(self.db,s.id)

    def test_profit_loss_filters_expenses_to_driver(self):
        self.load()
        self.db.add_all([Expense(category='Other',driver_id=self.driver.id,expense_date=TODAY,amount=100,is_active=True),
                         Expense(category='Other',expense_date=TODAY,amount=999,is_active=True)])
        self.db.commit()
        report=reports.get_profit_loss_report(self.db,date_from=TODAY,date_to=TODAY,driver_id=self.driver.id)
        self.assertEqual(report['summary']['expenses'],100)

    def test_lumper_deduction_has_negative_revenue(self):
        l=self.load();self.db.add(LoadService(load_id=l.id,service_type=ServiceType.LUMPER,add_deduct='Deduct',invoice_amount=100,drivers_payable=0))
        self.db.commit();self.db.expire(l,['services'])
        self.assertEqual(reports._row(l)['lumpers'],-100)

    def test_legacy_percentage_name_is_normalized(self):
        self.profile.pay_type='freight_percentage'; self.db.commit()
        l=self.load()
        self.assertEqual(l.pay_type_snapshot,'percentage')
        self.assertEqual(l.drivers_payable_snapshot,1800)

    def test_carryover_reversal_restores_both_sides(self):
        s=self.settlement();payroll.add_adjustment(self.db,s.id,SettlementAdjustmentCreate(adj_type='deduction',amount=200))
        self.ready(s);entry=payroll.create_carryover(self.db,s.id,TODAY);target=self.settlement()
        s.status=SettlementStatus.PREPARING;self.db.commit()
        payroll.remove_carryover(self.db,target.id,entry.id)
        self.assertEqual(s.balance_due,-200)
        self.assertEqual(target.balance_due,0)

    def test_new_settlement_cannot_start_paid(self):
        with self.assertRaises(payroll.PayrollError):
            payroll.create_settlement(self.db,SettlementCreate(driver_id=self.driver.id,date=TODAY,status=SettlementStatus.PAID))

    def test_pickup_filter_uses_stop_date_not_load_creation_date(self):
        l=self.load(load_date=TODAY-timedelta(days=5))
        self.db.add(LoadStop(load_id=l.id,stop_type=StopType.PICKUP,stop_order=1,stop_date=TODAY))
        self.db.commit()
        balances=payroll.get_open_balances(self.db,date_from=TODAY,date_to=TODAY,date_type='pickup')
        self.assertEqual(balances[0]['balance'],1200)
        report=reports.get_total_revenue_report(self.db,date_from=TODAY,date_to=TODAY,date_type='pickup')
        self.assertEqual(report['summary']['total_loads'],1)
        self.assertEqual(payroll.get_open_balances(self.db,date_from=TODAY,date_to=TODAY,date_type='delivery'),[])

    def test_legacy_advance_is_included_and_persisted(self):
        from app.api.v1.endpoints.payments import apply_to_settlement, unapply_from_settlement
        l=self.load();s=self.settlement();payroll.add_load_item(self.db,s.id,l.id)
        p=Payment(payment_number=4001,driver_id=self.driver.id,payment_type='advanced_payment',payment_date=TODAY,amount=100,is_active=True)
        self.db.add(p);self.db.commit()
        apply_to_settlement(p.id,s.id,self.db)
        self.db.expire_all()
        self.assertEqual((s.settlement_total,s.balance_due),(1200,1100))
        unapply_from_settlement(p.id,self.db)
        self.db.expire_all()
        self.assertEqual(s.balance_due,1200)

    def test_schema_rejects_invalid_money(self):
        from pydantic import ValidationError
        for amount in [0,-1,float('nan'),float('inf')]:
            with self.assertRaises(ValidationError):SettlementPaymentCreate(amount=amount)
        with self.assertRaises(ValidationError):SettlementAdjustmentCreate(adj_type='advanced_payment',amount=10)

    def test_settlement_pdf_agrees_with_advance_totals(self):
        from app.services.pdf_service import generate_settlement_pdf
        from pypdf import PdfReader
        from io import BytesIO
        l=self.load();s=self.settlement();payroll.add_load_item(self.db,s.id,l.id)
        ap=AdvancedPayment(payment_number=3001,driver_id=self.driver.id,payment_date=TODAY,amount=300,applied_amount=0)
        self.db.add(ap);self.db.commit();payroll_api.apply_advanced_payment(s.id,ap.id,None,self.db)
        s=payroll.get_settlement(self.db,s.id)
        content=generate_settlement_pdf(s,db=self.db)
        text=' '.join(page.extract_text() for page in PdfReader(BytesIO(content)).pages)
        self.assertIn('$1,200.00',text)
        self.assertIn('Applied advance payments: $300.00',text)
        self.assertIn('Balance due: $900.00',text)

    def test_payee_change_does_not_move_old_loads(self):
        old=self.load()
        self.profile.payable_to='New company';self.db.commit()
        new=self.load()
        groups={row['payable_to']:row['balance'] for row in payroll.get_open_balances(self.db)}
        self.assertEqual(groups,{'Driver A':1200,'New company':1200})
        s=self.settlement()
        with self.assertRaises(payroll.PayrollError):payroll.add_load_item(self.db,s.id,old.id)
        payroll.add_load_item(self.db,s.id,new.id)

    def test_settled_load_cannot_be_deleted(self):
        l=self.load();s=self.settlement();payroll.add_load_item(self.db,s.id,l.id)
        with self.assertRaises(ValueError):loads.delete_load(self.db,l.id)
        self.assertTrue(l.is_active)

    def test_reference_extra_stop_and_override_totals(self):
        from app.services.driver_pay_service import refresh_extra_stop_count
        self.driver.pay_rate_empty=.3; self.profile.per_extra_stop=25
        l=self.load(rate=1000, loaded_miles=100, empty_miles=10, stops=[
            LoadStop(stop_type=StopType.PICKUP, stop_order=1), LoadStop(stop_type=StopType.DELIVERY, stop_order=2)])
        self.assertEqual(compute_driver_pay(l),63)
        l.stops.append(LoadStop(stop_type=StopType.PICKUP, stop_order=3, is_payable=False))
        refresh_extra_stop_count(l); self.assertEqual(compute_driver_pay(l),88)
        l.stops[-1].is_payable=True
        refresh_extra_stop_count(l); self.assertEqual(compute_driver_pay(l),88)
        l.stops.append(LoadStop(stop_type=StopType.OTHER, stop_order=4, is_payable=False))
        refresh_extra_stop_count(l); self.assertEqual(compute_driver_pay(l),88)
        l.stops[-1].is_payable=True
        refresh_extra_stop_count(l); self.assertEqual(compute_driver_pay(l),113)
        l.driver_pay_override={'type':'fixed','amount':200}
        self.assertEqual(compute_driver_pay(l),250)
        l.driver_pay_override={'type':'percentage','base':1000,'percentage':20}
        self.assertEqual(compute_driver_pay(l),250)
        l.driver_pay_override={'type':'per_mile','loaded_rate':.5,'empty_rate':0,'extra_stop_rate':10}
        self.assertEqual(compute_driver_pay(l),70)
        self.profile.per_extra_stop=999; self.db.commit()
        l.driver_pay_override=None
        self.assertEqual(compute_driver_pay(l),113)

    def test_reference_payee_and_quickpay_different_bases(self):
        from app.models.models import Vendor, DriverAdditionalPayee, Broker
        from app.services.load_financials import capture_quickpay, quickpay_fee, invoice_amount, refresh_payees
        vendor=Vendor(company_name='Additional vendor',is_additional_payee=True)
        broker=Broker(name='Customer',quickpay_fee=2.25)
        self.db.add_all([vendor,broker]); self.db.flush()
        l=self.load(rate=1000,broker_id=broker.id)
        capture_quickpay(self.db,l)
        rule=DriverAdditionalPayee(driver_id=self.driver.id,vendor_id=vendor.id,rate_pct=10)
        self.db.add(rule);self.db.commit()
        self.assertEqual(l.additional_payees,[])
        loads.update_load(self.db,l.id,LoadUpdate(rate=1200))
        self.assertEqual(l.additional_payees[0].amount,120)
        self.assertEqual(quickpay_fee(l),27)
        l.services.append(LoadService(service_type=ServiceType.OTHER,invoice_amount=100,drivers_payable=0,add_deduct='Add'))
        self.assertEqual(invoice_amount(l),1300)
        self.assertEqual(quickpay_fee(l),29.25)
        self.assertEqual(l.additional_payees[0].amount,120)
        rule.rate_pct=50;broker.quickpay_fee=20;self.db.commit()
        refresh_payees(self.db,l)
        self.assertEqual(l.additional_payees[0].amount,120)
        self.assertEqual(quickpay_fee(l),29.25)

    def test_additional_payee_application_is_unique_reversible_and_locks_load(self):
        from app.models.models import Vendor, DriverAdditionalPayee
        from fastapi import HTTPException
        vendor=Vendor(company_name='Additional vendor',is_additional_payee=True)
        self.db.add(vendor);self.db.flush()
        self.db.add(DriverAdditionalPayee(driver_id=self.driver.id,vendor_id=vendor.id,rate_pct=10));self.db.commit()
        l=self.load(rate=1200);entry=l.additional_payees[0]
        wrong=self.settlement()
        with self.assertRaises(HTTPException): payroll_api.apply_additional_payee(wrong.id,entry.id,self.db)
        s=payroll.create_settlement(self.db,SettlementCreate(driver_id=self.driver.id,date=TODAY,payable_to=vendor.company_name))
        adj=payroll_api.apply_additional_payee(s.id,entry.id,self.db)
        self.assertEqual(payroll_api.apply_additional_payee(s.id,entry.id,self.db)['id'],adj['id'])
        self.assertEqual(s.settlement_total,120)
        self.assertEqual(payroll_api.available_additional_payees(s.id,self.db),[])
        with self.assertRaises(ValueError): loads.update_load(self.db,l.id,LoadUpdate(rate=1300))
        with self.assertRaises(ValueError): loads.delete_load(self.db,l.id)
        payroll.delete_adjustment(self.db,s.id,adj['id'])
        self.assertEqual(len(payroll_api.available_additional_payees(s.id,self.db)),1)
        payroll_api.apply_additional_payee(s.id,entry.id,self.db)
        payroll.delete_settlement(self.db,s.id)
        self.assertEqual(self.db.query(SettlementAdjustment).filter(SettlementAdjustment.load_payee_id==entry.id).count(),0)

    def test_frontend_api_paths_match_registered_routes(self):
        from app.api.v1 import api_router
        from pathlib import Path
        routes={r.path for r in api_router.routes}
        self.assertIn('/api/v1/payroll/{settlement_id}/time-reports',routes)
        self.assertIn('/api/v1/loads/{load_id}/driver-pay-override',routes)
        self.assertIn('/api/v1/scheduled-transactions/preview',routes)
        self.assertIn('/api/v1/drivers/{driver_id}/additional-payees',routes)
        # Check the actual UI request strings against the public API mount.
        root=Path(__file__).resolve().parents[2]/'frontend/src'
        for file,path in [('components/payroll/SettlementModal.tsx','/api/v1/payroll/${settlementId}/time-reports'),
                          ('components/loads/LoadModal.tsx','/api/v1/loads/${load.id}/driver-pay-override'),
                          ('pages/DriversPage.tsx','/api/v1/scheduled-transactions/preview')]:
            self.assertIn(path,(root/file).read_text())

    def test_http_payroll_and_override_contract(self):
        import asyncio, httpx
        from fastapi import FastAPI
        from app.api.v1 import api_router
        from app.db.session import get_db
        app=FastAPI();app.include_router(api_router)
        def test_db(): yield self.db
        app.dependency_overrides[get_db]=test_db
        l=self.load();settlement=self.settlement()
        async def requests():
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app),base_url='http://test') as c:
                response=await c.get(f'/api/v1/payroll/{settlement.id}/time-reports')
                self.assertEqual(response.status_code,200,response.text)
                response=await c.put(f'/api/v1/loads/{l.id}/driver-pay-override',json={'type':'fixed','amount':200})
                self.assertEqual(response.status_code,200,response.text)
                self.assertEqual(self.db.get(Load,l.id).drivers_payable_snapshot,200)
                response=await c.get(f'/api/v1/drivers/{self.driver.id}/additional-payees')
                self.assertEqual(response.status_code,200,response.text)
                response=await c.get(f'/api/v1/payroll/{settlement.id}/additional-payees')
                self.assertEqual(response.status_code,200,response.text)
        asyncio.run(requests())

    @patch("app.services.report_export_service._get_company_info", return_value={"name":"KARVAN Test","email":"","phone":"","address":""})
    def test_reference_fees_reconcile_reports_and_exports(self, _company):
        import json, io
        from pathlib import Path
        from openpyxl import load_workbook
        from pypdf import PdfReader
        from app.models.models import Vendor, DriverAdditionalPayee
        from app.services.report_export_service import generate_gross_profit_per_load_pdf, generate_gross_profit_per_load_xlsx
        fixture=json.loads((Path(__file__).parent/'fixtures/ezloads_load_financials.json').read_text())['fee_accessorial']
        vendor=Vendor(company_name='Additional vendor',is_additional_payee=True)
        self.db.add(vendor);self.db.flush()
        self.db.add(DriverAdditionalPayee(driver_id=self.driver.id,vendor_id=vendor.id,rate_pct=10));self.db.commit()
        l=self.load(rate=fixture['amount_freight'],loaded_miles=100,empty_miles=0)
        l.quickpay_rate_snapshot=2.25;l.drivers_payable_snapshot=70
        l.services.append(LoadService(service_type=ServiceType.OTHER,invoice_amount=100,drivers_payable=0,add_deduct='Add'))
        self.db.commit()
        data=reports.get_gross_profit_per_load_report(self.db,date_from=TODAY,date_to=TODAY)
        row=data['rows'][0]
        self.assertEqual(row['additional_payee'],float(fixture['amount_additional_payee']))
        self.assertEqual(row['qp_fee'],fixture['amount_funding_fees'])
        self.assertEqual(row['gross_profit'],1080.75)
        pdf=generate_gross_profit_per_load_pdf(data,{})
        text=' '.join(page.extract_text() for page in PdfReader(io.BytesIO(pdf)).pages)
        self.assertIn('Additional Payees',' '.join(text.split()))
        self.assertIn('1,080.75',text)
        workbook=load_workbook(io.BytesIO(generate_gross_profit_per_load_xlsx(data,{})))
        rows=list(workbook.active.values)
        headers=next(r for r in rows if 'Additional Payees' in r)
        total=next(r for r in rows if 'Total' in r)
        self.assertEqual(total[headers.index('Additional Payees')],120)
        self.assertEqual(total[headers.index('Gross Profit')],1080.75)
        # Scratch output for visual review; not an application or production write.
        Path('/tmp/karvan-payee-report.pdf').write_bytes(pdf)

    def test_deploy_startup_preserves_existing_accounts(self):
        import ast, traceback
        from pathlib import Path
        from app.models.models import User
        user=User(name='Existing name',email='admin@karvan.com',hashed_password='existing-hash',role='dispatcher',is_active=False)
        self.db.add(user);self.db.commit()
        source=ast.parse((Path(__file__).parents[1]/'app/main.py').read_text())
        function=next(n for n in source.body if isinstance(n,ast.FunctionDef) and n.name=='startup_fix_snapshots')
        function.decorator_list=[]
        namespace={'traceback':traceback}
        exec(compile(ast.fix_missing_locations(ast.Module(body=[function],type_ignores=[])),'startup-test','exec'),namespace)
        with patch('app.db.session.SessionLocal',return_value=self.db), patch('app.services.auth_service.hash_password',return_value='new-hash'):
            namespace['startup_fix_snapshots']()
        saved=self.db.query(User).filter_by(email='admin@karvan.com').one()
        self.assertEqual(saved.hashed_password,'existing-hash')
        self.assertEqual(saved.name,'Existing name')
        self.assertEqual(saved.role.value,'dispatcher')
        self.assertFalse(saved.is_active)

if __name__ == '__main__': unittest.main()
